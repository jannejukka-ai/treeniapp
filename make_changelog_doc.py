#!/usr/bin/env python3
"""
make_changelog_doc.py — tekee CHANGELOG.md:stä muotoillun Word-dokumentin.

Käyttö:
    python3 make_changelog_doc.py

Lukee samasta kansiosta CHANGELOG.md ja kirjoittaa Treeniapp-Versiohistoria.docx.

Muotoilu:
    ## v3.2 — ...   -> sininen versio-otsikko
    ### Alaotsikko   -> tumma alaotsikko
    **Lihavoitu:**   -> tumma väliotsikko (rivi joka on kokonaan lihavoitu)
    - kohta          -> bulletti (myös sisennetyt)
    1. kohta         -> numeroitu kohta
    muu teksti       -> leipäteksti
    **lihava** tekstin seassa säilyy lihavana.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "CHANGELOG.md"
OUTPUT = HERE / "Treeniapp-Versiohistoria.docx"

BLUE = RGBColor(0x1D, 0x4E, 0xD8)   # versio-otsikko
DARK = RGBColor(0x1E, 0x29, 0x3B)   # alaotsikot
GREY = RGBColor(0x64, 0x74, 0x8B)   # aputeksti


def add_rich_text(paragraph, text, base_bold=False, color=None, size=None):
    """Lisää tekstin kappaleeseen niin että **lihavointi** säilyy."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part.replace("`", ""))
        run.bold = base_bold or (i % 2 == 1)
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = Pt(size)


def build(source: Path, output: Path) -> Path:
    if not source.exists():
        sys.exit(f"Ei löydy: {source}")

    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich_text(title, "Treeniapp — Versiohistoria", base_bold=True, color=BLUE, size=22)

    sub = doc.add_paragraph()
    add_rich_text(sub, "Automaattisesti koottu CHANGELOG.md-tiedostosta.", color=GREY, size=9)

    skip_intro = True

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        # Ohita tiedoston oma otsikko ja johdanto ensimmäiseen versioon asti
        if skip_intro:
            if stripped.startswith("## "):
                skip_intro = False
            else:
                continue

        if not stripped or stripped == "---":
            continue

        # Versio-otsikko
        if stripped.startswith("## "):
            doc.add_paragraph()
            p = doc.add_paragraph()
            add_rich_text(p, stripped[3:], base_bold=True, color=BLUE, size=16)
            continue

        # Alaotsikko
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            add_rich_text(p, stripped[4:], base_bold=True, color=DARK, size=12)
            continue

        # Kokonaan lihavoitu rivi = väliotsikko
        if re.fullmatch(r"\*\*.+\*\*:?", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            add_rich_text(p, stripped, color=DARK, size=11)
            continue

        # Bullet (myös sisennetty)
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent, text = m.group(1), m.group(2)
            style_name = "List Bullet 2" if len(indent) >= 2 else "List Bullet"
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, text)
            continue

        # Numeroitu kohta
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, m.group(1))
            continue

        # Leipäteksti
        p = doc.add_paragraph()
        add_rich_text(p, stripped)

    doc.save(output)
    return output


if __name__ == "__main__":
    path = build(SOURCE, OUTPUT)
    print(f"Valmis: {path}")
